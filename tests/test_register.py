import datetime
import json

import openpyxl  # dev-only independent reader; the tool itself uses the standard library
import pytest

import register as reg_tool


@pytest.fixture
def folder(tmp_path):
    reg_tool.save(tmp_path, reg_tool.new_register("Cedar"))
    return tmp_path


def change(base, op, *changes, **extra):
    return {"base_revision": base, "op_id": op, "actor": "test", "summary": op, "changes": list(changes), **extra}


RISK = {"sheet": "risks", "add": {"title": "FX mismatch", "raised_by": "E&S specialist", "rating": "high",
                                  "destination": "main", "reason": "affects repayment", "status": "open"}}


def test_ids_are_assigned_and_revision_advances(folder):
    r = reg_tool.apply(folder, change(0, "a", RISK, {"sheet": "risks", "add": dict(RISK["add"], title="Customer")}))
    assert r == {"status": "applied", "revision": 1, "ids": ["R-001", "R-002"], "warnings": []}
    assert reg_tool.load(folder)["log"][-1]["op_id"] == "a"


def test_stale_and_duplicate_change_sets(folder):
    reg_tool.apply(folder, change(0, "a", RISK))
    with pytest.raises(reg_tool.ChangeError, match="stale"):
        reg_tool.apply(folder, change(0, "b", RISK))
    assert reg_tool.apply(folder, change(0, "a", RISK))["status"] == "already-applied"
    assert len(reg_tool.load(folder)["risks"]) == 1


@pytest.mark.parametrize("bad, message", [
    ({"sheet": "mitigants", "add": {"risk": "R-009", "raised_by": "x"}}, "unknown risk"),
    ({"sheet": "mitigants", "add": {"risk": "R-001", "raised_by": "x", "verified": "yes"}}, "needs a source"),
    ({"sheet": "open_items", "add": {"text": "q", "raised_by": "x", "status": "answered"}}, "answer source"),
    ({"sheet": "risks", "add": {"title": "t", "rating": "high"}}, "who raised"),
    ({"sheet": "risks", "add": {"title": "t", "raised_by": "x", "rating": "huge"}}, "rating"),
    ({"sheet": "risks", "update": "R-001", "set": {"destination": "background"}}, "officer's decision"),
    ({"sheet": "open_items", "add": {"text": "q", "raised_by": "x", "link": "R-001", "criticality": "later"}},
     "officer_ok"),
    ({"sheet": "nonsense", "add": {}}, "unknown sheet"),
    ({"sheet": "risks", "add": {"colour": "red"}}, "no field"),
])
def test_rule_violations_reject_the_whole_change_set(folder, bad, message):
    reg_tool.apply(folder, change(0, "a", RISK))
    with pytest.raises(reg_tool.ChangeError, match=message):
        reg_tool.apply(folder, change(1, "b", {"sheet": "sources", "add": {"title": "kept?"}}, bad))
    assert reg_tool.load(folder)["sources"] == []   # nothing from the rejected set was saved


def test_downgrade_needs_officer_decision(folder):
    reg_tool.apply(folder, change(0, "a", RISK))
    with pytest.raises(reg_tool.ChangeError, match="downgrading high to medium"):
        reg_tool.apply(folder, change(1, "b", {"sheet": "risks", "update": "R-001", "set": {"rating": "medium"}}))
    ok = {"sheet": "risks", "update": "R-001",
          "set": {"rating": "medium", "officer_decision": "Officer, 2 Oct: hedge confirmed"}}
    assert reg_tool.apply(folder, change(1, "c", ok))["status"] == "applied"


def test_officer_ok_allows_later_diligence(folder):
    reg_tool.apply(folder, change(0, "a", RISK))
    item = {"text": "Check FX hedging policy", "raised_by": "credit", "link": "R-001", "criticality": "later",
            "officer_ok": "yes", "status": "open"}
    assert reg_tool.apply(folder, change(1, "b", {"sheet": "open_items", "add": item}))["status"] == "applied"


def test_ready_gate_blocks_and_word_change_invalidates(folder, tmp_path):
    word = tmp_path / "FP-v01.docx"
    word.write_bytes(b"v1")
    reg_tool.apply(folder, change(0, "a", RISK, {"sheet": "open_items", "add": {
        "text": "Borrower cash flow", "raised_by": "officer", "criticality": "blocker", "status": "open",
        "addressee": "client"}}, {"sheet": "control", "set": {"word": word.name, "word_sha256": reg_tool.sha256(word)}}))
    with pytest.raises(reg_tool.ChangeError, match="cannot record ready"):
        reg_tool.apply(folder, change(1, "b", {"sheet": "control", "set": {"status": "ready"}}))
    reg_tool.apply(folder, change(1, "c", {"sheet": "open_items", "update": "O-001",
                                           "set": {"status": "closed", "answer_source": "S-004 p.2"}}))
    reg = reg_tool.load(folder)
    assert reg_tool.open_checks(reg, word) == []
    word.write_bytes(b"v1 edited by officer")
    assert reg_tool.open_checks(reg, word) == ["the Word file changed since it was last recorded"]


def test_new_source_with_pending_impact_blocks_ready(folder):
    reg_tool.apply(folder, change(0, "a", {"sheet": "sources", "add": {"title": "Lender email: covenant breach",
                                                                       "impact": "pending", "tier": "A"}}))
    reasons = reg_tool.open_checks(reg_tool.load(folder))
    assert reasons == ["source impact not assessed: S-001"]


def test_stale_sources_warning(folder):
    old = (datetime.date.today() - datetime.timedelta(days=500)).isoformat()
    reg_tool.apply(folder, change(0, "a", {"sheet": "sources", "add": {"title": "AR 2023", "event_date": old}},
                                  {"sheet": "fp_map", "add": {"version": "v01", "section": "summary", "ids": "S-001"}}))
    assert any(p[1] == "stale-sources" for p in reg_tool.problems(reg_tool.load(folder)))


def test_analysis_upsert_and_markdown_view(folder):
    reg_tool.apply(folder, change(0, "a", {"sheet": "analysis", "update": "timeline",
                                           "set": {"title": "Last two years", "text": "2025: new plant."}}))
    reg_tool.apply(folder, change(1, "b", {"sheet": "analysis", "update": "timeline", "set": {"text": "2026: CFO left."}}))
    md = (folder / "Analysis.md").read_text()
    assert "2026: CFO left." in md and "2025: new plant." not in md and "revision 2" in md


def test_excel_view_opens_in_an_independent_reader(folder):
    reg_tool.apply(folder, change(0, "a", RISK))
    wb = openpyxl.load_workbook(folder / "register.xlsx")
    assert wb.sheetnames[:3] == ["Control", "Sources", "Risks"]
    assert wb["Risks"]["A2"].value == "R-001" and wb["Risks"]["B2"].value == "FX mismatch"
    assert wb["Control"]["B2"].value == "Cedar"


def test_summary_lists_show_and_cli(folder, tmp_path, capsys):
    reg_tool.apply(folder, change(0, "a", RISK, {"sheet": "open_items", "add": {
        "text": "Revenue split by currency", "raised_by": "officer", "link": "R-001", "criticality": "blocker",
        "status": "open", "addressee": "client", "due": "2026-10-10"}},
        {"sheet": "mitigants", "add": {"risk": "R-001", "text": "Hedging covenant", "type": "proposed-condition",
                                       "raised_by": "credit"}}, next_step="Grill v01 in a new task"))
    reg = reg_tool.load(folder)
    text = reg_tool.summary(reg)
    assert "Next: Grill v01 in a new task" in text and "O-001 [blocker] client" in text
    assert len(text.splitlines()) <= 40
    assert "Revenue split by currency" in reg_tool.lists(reg) and "M-001 (R-001)" in reg_tool.lists(reg)
    assert "mitigant M-001" in reg_tool.show(reg, ["R-001"])
    changes = tmp_path / "c.json"
    changes.write_text(json.dumps(change(1, "cli", {"sheet": "control", "set": {"officer": "Rik"}})))
    assert reg_tool.main(["apply", str(folder), str(changes)]) == 0
    assert reg_tool.main(["apply", str(folder), str(changes)]) == 0   # retry is a no-op
    assert reg_tool.load(folder)["revision"] == 2
    assert reg_tool.main(["selftest"]) == 0


def test_new_questions_default_open_and_new_sources_pending(folder):
    reg_tool.apply(folder, change(0, "a", {"sheet": "open_items", "add": {"text": "Borrower forecast", "raised_by": "officer"}},
                                  {"sheet": "sources", "add": {"title": "Email from CFO"}}))
    reg = reg_tool.load(folder)
    assert reg["open_items"][0]["status"] == "open" and reg["sources"][0]["impact"] == "pending"
    assert "O-001" in reg_tool.lists(reg)


def test_promised_and_answered_items_stay_visible_until_closed(folder):
    reg_tool.apply(folder, change(0, "a", {"sheet": "open_items", "add": {
        "text": "Monthly borrower forecast", "raised_by": "officer", "criticality": "blocker", "addressee": "client"}}))
    # "We will send the forecast on Friday" is a promise, not the forecast.
    reg_tool.apply(folder, change(1, "b", {"sheet": "open_items", "update": "O-001", "set": {
        "status": "promised", "answer": "CFO: will send Friday", "due": "2026-10-09"}}))
    reg = reg_tool.load(folder)
    assert "O-001" in reg_tool.lists(reg) and any("O-001" in r for r in reg_tool.gates(reg))
    reg_tool.apply(folder, change(2, "c", {"sheet": "open_items", "update": "O-001", "set": {
        "status": "answered", "answer_source": "S-001 forecast file"}}))
    reg = reg_tool.load(folder)
    assert "awaiting assessment" in reg_tool.lists(reg) and any("O-001" in r for r in reg_tool.gates(reg))
    client_requests, internal = reg_tool.lists(reg).split('## Internal open points')
    assert 'O-001' not in client_requests and 'O-001' in internal
    reg_tool.apply(folder, change(3, "d", {"sheet": "open_items", "update": "O-001", "set": {"status": "closed"}}))
    reg = reg_tool.load(folder)
    assert "O-001" not in reg_tool.lists(reg) and reg_tool.gates(reg) == []


def test_tool_reports_checks_and_readiness_is_recorded_for_a_version(folder, tmp_path):
    word = tmp_path / "FP-v05.docx"
    word.write_bytes(b"v5")
    reg_tool.apply(folder, change(0, "a", {"sheet": "control", "set": {"word": word.name, "word_sha256": reg_tool.sha256(word)}}))
    assert "Register checks: passed" in reg_tool.summary(reg_tool.load(folder), word)
    reg_tool.apply(folder, change(1, "b", {"sheet": "control", "set": {"status": "ready", "ready_by": "Officer"}}))
    reg = reg_tool.load(folder)
    assert reg["control"]["ready_for"] == f"{word.name} at register revision 2"
    assert "Status: ready (FP-v05.docx at register revision 2, recorded by Officer)" in reg_tool.summary(reg, word)
    word.write_bytes(b"v5 edited")
    assert "changed since readiness was recorded" in reg_tool.summary(reg, word)


def test_retry_repairs_views_left_stale_by_an_interrupted_run(folder, monkeypatch):
    real = reg_tool.write_views
    monkeypatch.setattr(reg_tool, "write_views", lambda *a: (_ for _ in ()).throw(OSError("disk full")))
    with pytest.raises(OSError):
        reg_tool.apply(folder, change(0, "a", RISK))
    monkeypatch.setattr(reg_tool, "write_views", real)
    assert "revision 0" in (folder / "Analysis.md").read_text()      # views are stale, the record is not
    assert reg_tool.apply(folder, change(0, "a", RISK))["status"] == "already-applied"
    assert "revision 1" in (folder / "Analysis.md").read_text()


def test_new_evidence_after_ready_withdraws_readiness_instead_of_blocking(folder):
    record_word(folder)
    reg_tool.apply(folder, change(1, "a", {"sheet": "control", "set": {"status": "ready", "ready_by": "FP Lead"}}))
    r = reg_tool.apply(folder, change(2, "b", {"sheet": "sources", "add": {"title": "Lender email: covenant breach"}}))
    assert r["status"] == "applied" and r["warnings"][0].startswith("readiness withdrawn")
    assert reg_tool.load(folder)["control"]["status"] == "draft"


def record_word(folder, revision=0):
    word = folder / 'FP.docx'
    word.write_bytes(b'current synthetic Word version')
    reg_tool.apply(folder, change(revision, 'word', {'sheet': 'control', 'set': {
        'word': word.name, 'word_sha256': reg_tool.sha256(word)}}))
    return word


def test_findings_default_open_and_block_readiness(folder):
    reg_tool.apply(folder, change(0, 'finding', {'sheet': 'findings', 'add': {
        'by': 'grill', 'issue': 'Unsupported conclusion', 'priority': 'blocker'}}))
    reg = reg_tool.load(folder)
    assert reg['findings'][0]['status'] == 'open'
    assert reg_tool.gates(reg) == ['open blockers: G-001']


@pytest.mark.parametrize('sheet, fields, field', [
    ('sources', {'title': 'Evidence'}, 'impact'),
    ('sources', {'title': 'Evidence'}, 'coverage'),
    ('risks', RISK['add'], 'status'),
    ('open_items', {'text': 'Question', 'raised_by': 'officer', 'criticality': 'blocker'}, 'status'),
    ('open_items', {'text': 'Question', 'raised_by': 'officer', 'criticality': 'blocker'}, 'criticality'),
    ('findings', {'by': 'grill', 'priority': 'blocker'}, 'status'),
    ('findings', {'by': 'grill', 'priority': 'blocker'}, 'priority'),
])
def test_lifecycle_or_blocker_fields_cannot_be_cleared(folder, sheet, fields, field):
    result = reg_tool.apply(folder, change(0, 'add', {'sheet': sheet, 'add': fields}))
    with pytest.raises(reg_tool.ChangeError, match=field):
        reg_tool.apply(folder, change(1, 'clear', {'sheet': sheet, 'update': result['ids'][0], 'set': {field: ''}}))
    assert reg_tool.load(folder)['revision'] == 1


@pytest.mark.parametrize('fields', [
    {'status': 'ready'},
    {'status': 'ready', 'ready_by': '  '},
    {'status': 'ready', 'ready_by': 'Officer', 'word': '', 'word_sha256': ''},
    {'status': 'ready', 'ready_by': 'Officer', 'word_sha256': 'old'},
    {'status': 'ready', 'ready_by': 'Officer', 'word': 'missing.docx'},
])
def test_ready_requires_explicit_officer_and_current_word(folder, fields):
    record_word(folder)
    with pytest.raises(reg_tool.ChangeError, match='cannot record ready'):
        reg_tool.apply(folder, change(1, 'ready', {'sheet': 'control', 'set': fields}, actor='Automation'))
    assert reg_tool.load(folder)['control']['status'] == 'draft'


@pytest.mark.parametrize('edit', ['evidence', 'word_metadata', 'word_bytes'])
def test_ready_invalidated_on_assessed_evidence_or_word_changes(folder, edit):
    word = record_word(folder)
    reg_tool.apply(folder, change(1, 'ready', {'sheet': 'control', 'set': {'status': 'ready', 'ready_by': 'Officer'}}))
    changes = []
    if edit == 'evidence':
        changes = [{'sheet': 'sources', 'add': {'title': 'New facts', 'impact': 'assessed'}}]
    else:
        word.write_bytes(b'changed Word')
        if edit == 'word_metadata':
            changes = [{'sheet': 'control', 'set': {'word_sha256': reg_tool.sha256(word)}}]
        else:
            text = reg_tool.summary(reg_tool.load(folder), word)
            assert 'Status: ready' not in text
            assert 'changed since readiness was recorded' in text
    result = reg_tool.apply(folder, change(2, 'changed', *changes))
    assert result['warnings'][0].startswith('readiness withdrawn')
    assert reg_tool.load(folder)['control']['status'] == 'draft'


def test_views_have_revision_and_partial_save_is_honest(folder, monkeypatch, capsys):
    changes = folder / 'changes.json'
    changes.write_text(json.dumps(change(0, 'save', RISK)))
    original = reg_tool.write_views
    monkeypatch.setattr(reg_tool, 'write_views', lambda *a: (_ for _ in ()).throw(OSError('disk full')))
    assert reg_tool.main(['apply', str(folder), str(changes)]) == 4
    err = capsys.readouterr().err
    assert 'saved' in err.lower() and 'views incomplete' in err.lower() and 'REFUSED' not in err
    monkeypatch.setattr(reg_tool, 'write_views', original)
    assert reg_tool.main(['apply', str(folder), str(changes)]) == 0
    wb = openpyxl.load_workbook(folder / 'register.xlsx')
    control = dict(wb['Control'].values)
    assert control['register_revision'] == 1
    assert 'revision 1' in (folder / 'Analysis.md').read_text()


@pytest.mark.parametrize('bad, message', [
    ({'sheet': 'risks', 'add': {'title': 'risk', 'raised_by': '  '}}, 'who raised'),
    ({'sheet': 'mitigants', 'add': {'risk': 'R-001', 'raised_by': 'x', 'verified': 'yes', 'source': '  '}}, 'needs a source'),
    ({'sheet': 'open_items', 'add': {'text': 'q', 'raised_by': 'x', 'status': 'closed', 'answer_source': '  '}}, 'answer source'),
    ({'sheet': 'risks', 'update': 'R-001', 'set': {'rating': 'medium', 'officer_decision': '  '}}, "officer's decision"),
])
def test_blank_required_evidence_and_decisions_are_refused(folder, bad, message):
    reg_tool.apply(folder, change(0, 'risk', RISK))
    with pytest.raises(reg_tool.ChangeError, match=message):
        reg_tool.apply(folder, change(1, 'bad', bad))
    assert reg_tool.load(folder)['revision'] == 1


def test_blocker_awaiting_officer_judgment_still_blocks_readiness(folder):
    record_word(folder)
    reg_tool.apply(folder, change(1, 'finding', {'sheet': 'findings', 'add': {
        'by': 'grill', 'issue': 'Repayment depends on unsigned refinancing',
        'priority': 'blocker', 'status': 'officer-judgment'}}))
    assert 'G-001' in '; '.join(reg_tool.gates(reg_tool.load(folder)))
    assert 'Findings: 1 unresolved' in reg_tool.summary(reg_tool.load(folder))
    with pytest.raises(reg_tool.ChangeError, match='open blockers: G-001'):
        reg_tool.apply(folder, change(2, 'ready', {'sheet': 'control', 'set': {
            'status': 'ready', 'ready_by': 'Officer'}}))


@pytest.mark.parametrize('payload', [[], {'changes': None}, {'changes': [None]},
    {'changes': [{'sheet': 'sources', 'add': None}]},
    {'changes': [{'sheet': 'sources', 'add': {'title': None}}]}])
def test_malformed_change_sets_are_refused_without_writing(folder, payload):
    if isinstance(payload, dict):
        payload = dict(base_revision=0, op_id='bad', **payload)
    with pytest.raises(reg_tool.ChangeError):
        reg_tool.apply(folder, payload)
    assert reg_tool.load(folder)['revision'] == 0


def test_risk_analysis_does_not_hide_risk_and_linked_items(folder):
    reg_tool.apply(folder, change(0, 'linked', RISK,
        {'sheet': 'analysis', 'add': {'id': 'R-001', 'text': 'Detailed FX exposure.'}},
        {'sheet': 'open_items', 'add': {'text': 'Confirm hedge', 'link': 'R-001', 'raised_by': 'officer'}}))
    shown = reg_tool.show(reg_tool.load(folder), ['R-001'])
    assert 'R-001 (risks)' in shown
    assert 'Confirm hedge' in shown and 'Detailed FX exposure.' in shown
