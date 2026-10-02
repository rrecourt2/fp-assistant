import json
import re
import shutil
import zipfile
from pathlib import Path

import docx  # dev-only independent reader; the tool itself uses the standard library
import pytest

import fp_docx


@pytest.fixture
def tpl(tmp_path):
    path = tmp_path / "template.docx"
    fp_docx.make_template(path)
    return path


def write(path, text):
    path.write_text(text, encoding="utf-8")
    return path


FIRST = ("## summary\nCedar Foods seeks EUR 12m. [S-001 p.2]\n\n"
         "## financial-analysis\nRevenue rose 20%. [S-002 p.20]\n\n| EUR m | FY2025 |\n|---|---|\n| Revenue | 120 |\n\n"
         "{{keep:1}}\n\n## market-risk\n- Milk prices rose. [S-003]\n\n## recommendation\nApprove.\n")


def test_first_build_fills_sections_fields_and_keeps_fixed_table(tpl, tmp_path):
    out = tmp_path / "FP-v01.docx"
    r = fp_docx.build(tpl, write(tmp_path / "c.md", FIRST), out, mode="first", template=tpl,
                      fields={"Borrower": "Cedar Foods", "Amount": "EUR 12m", "Tenor": "5 years"})
    assert r["written"] == ["summary", "financial-analysis", "market-risk", "recommendation"]
    assert r["missing_fields"] == ["Tenor"]
    d = docx.Document(out)  # opens in an independent reader
    text = "\n".join(p.text for p in d.paragraphs)
    assert "Cedar Foods seeks EUR 12m." in text and "[S-001" not in text
    assert "[Summarise the proposal]" not in text
    cells = [c.text for t in d.tables for row in t.rows for c in row.cells]
    assert "Cedar Foods" in cells and "EUR 12m" in cells   # cover fields
    assert "FY2025 (template)" in cells and "Revenue" in cells   # template table kept + new table added
    side = json.loads((tmp_path / "FP-v01.sections.json").read_text())
    assert side["sections"]["summary"]["owner"] == "ai"
    assert side["sections"]["summary"]["ids"] == ["S-001"]
    assert fp_docx.check(out, tpl)["problems"] == []


def test_second_round_freezes_officer_edits_and_writes_proposals(tpl, tmp_path):
    v1 = tmp_path / "FP-v01.docx"
    fp_docx.build(tpl, write(tmp_path / "c1.md", FIRST), v1, mode="first", template=tpl)
    fp_docx.edit_text(v1, "Approve.", "Approve, subject to two conditions.")
    v2 = tmp_path / "FP-v02.docx"
    r = fp_docx.build(v1, write(tmp_path / "c2.md", "## summary\nNew summary.\n\n## recommendation\nDecline.\n"), v2)
    assert r["written"] == ["summary"]
    assert r["frozen"] == {"recommendation": "edited since the assistant wrote it"}
    md = fp_docx.to_markdown(v2)
    assert "Approve, subject to two conditions." in md and "Decline." not in md
    assert "Decline." in (tmp_path / "FP-v02.proposals.md").read_text()
    side = json.loads((tmp_path / "FP-v02.sections.json").read_text())["sections"]
    assert side["recommendation"]["owner"] == "other" and side["market-risk"]["owner"] == "ai"


def test_commented_section_is_never_replaced(tmp_path):
    tpl = tmp_path / "template.docx"
    fp_docx.make_template(tpl, comment_on="market-risk")
    r = fp_docx.build(tpl, write(tmp_path / "c.md", FIRST), tmp_path / "out.docx", mode="first", template=tpl)
    assert "market-risk" in r["frozen"]
    info = fp_docx.inspect(tmp_path / "out.docx")
    assert info["comments"][0]["section"] == "market-risk"
    assert info["comments"][0]["text"] == "Check the FX exposure"


def test_unknown_section_goes_to_proposals(tpl, tmp_path):
    r = fp_docx.build(tpl, write(tmp_path / "c.md", "## annex-9\nText\n"), tmp_path / "o.docx",
                      mode="first", template=tpl)
    assert r["unmatched"] == ["annex-9"] and r["proposals"]


def test_refuses_to_overwrite(tpl, tmp_path):
    out = tmp_path / "o.docx"
    out.write_bytes(b"x")
    with pytest.raises(fp_docx.Refuse):
        fp_docx.build(tpl, write(tmp_path / "c.md", FIRST), out, mode="first", template=tpl)
    assert out.read_bytes() == b"x"


def test_fields_in_a_later_round_change_only_the_requested_cover_value(tpl, tmp_path):
    v1, v2 = tmp_path / "v1.docx", tmp_path / "v2.docx"
    fp_docx.build(tpl, write(tmp_path / "c.md", FIRST), v1, mode="first", template=tpl,
                  fields={"Borrower": "Cedar Foods", "Amount": "EUR 12m"})
    fp_docx.edit_text(v1, "Cedar Foods</w:t>", "Cedar Foods B.V.</w:t>")   # the officer's edit in Word
    r = fp_docx.build(v1, write(tmp_path / "c2.md", "## summary\nNew summary.\n"), v2, fields={"Amount": "EUR 15m"})
    cells = [c.text for t in docx.Document(v2).tables for row in t.rows for c in row.cells]
    assert "EUR 15m" in cells and "EUR 12m" not in cells
    assert "Cedar Foods B.V." in cells
    assert r["missing_fields"] == []


def test_signature_ignores_word_save_noise(tpl, tmp_path):
    v1 = tmp_path / "v1.docx"
    fp_docx.build(tpl, write(tmp_path / "c.md", FIRST), v1, mode="first", template=tpl)
    before = fp_docx.signatures(fp_docx.Doc(v1))
    with zipfile.ZipFile(v1) as z:
        parts = {i.filename: z.read(i.filename) for i in z.infolist()}
    xml = parts["word/document.xml"]
    # Word adds revision ids and splits runs when it saves; meaning is unchanged.
    xml = xml.replace(b"<w:p>", b'<w:p w:rsidR="00AB12CD">')
    xml = xml.replace(b'<w:t xml:space="preserve">Approve.</w:t>',
                      b'<w:t xml:space="preserve">Appr</w:t></w:r><w:r><w:t xml:space="preserve">ove.</w:t>')
    parts["word/document.xml"] = xml
    with zipfile.ZipFile(v1, "w") as z:
        for name, data in parts.items():
            z.writestr(name, data)
    after = fp_docx.signatures(fp_docx.Doc(v1))
    assert before == after


def test_bold_change_counts_as_an_edit(tpl, tmp_path):
    v1 = tmp_path / "v1.docx"
    fp_docx.build(tpl, write(tmp_path / "c.md", FIRST), v1, mode="first", template=tpl)
    fp_docx.edit_text(v1, '<w:r><w:t xml:space="preserve">Approve.</w:t>',
                      '<w:r><w:rPr><w:b/></w:rPr><w:t xml:space="preserve">Approve.</w:t>')
    assert fp_docx.inspect(v1)["sections"][-1]["changed_since_build"] is True


def test_cli_selftest_and_check_exit_codes(tpl, tmp_path, capsys):
    assert fp_docx.main(["selftest"]) == 0
    assert "PASS" in capsys.readouterr().out
    assert fp_docx.main(["check", str(tpl), "--template", str(tpl)]) == 1   # template still has guidance


def test_officer_edit_survives_a_round_that_touched_another_section(tpl, tmp_path):
    # Review finding: officer edits A in v1; v2 updates B only; v3 must still refuse to overwrite A.
    v1, v2, v3 = (tmp_path / f"v{n}.docx" for n in (1, 2, 3))
    fp_docx.build(tpl, write(tmp_path / "c1.md", FIRST), v1, mode="first", template=tpl)
    fp_docx.edit_text(v1, "Approve.", "Approve, in the officer's words.")
    fp_docx.build(v1, write(tmp_path / "c2.md", "## summary\nNew summary.\n"), v2)
    r = fp_docx.build(v2, write(tmp_path / "c3.md", "## recommendation\nDecline.\n"), v3)
    assert "recommendation" in r["frozen"]
    assert "officer's words" in fp_docx.to_markdown(v3)


def test_reviewer_can_read_kept_tables_and_reusing_the_text_does_not_duplicate_them(tpl, tmp_path):
    v1, v2 = tmp_path / "v1.docx", tmp_path / "v2.docx"
    fp_docx.build(tpl, write(tmp_path / "c.md", FIRST), v1, mode="first", template=tpl)
    md = fp_docx.to_markdown(v1)
    assert "FY2025 (template)" in md                     # the template's kept table is readable
    fp_docx.build(v1, write(tmp_path / "again.md", md), v2)
    assert len(docx.Document(v2).tables) == len(docx.Document(v1).tables)


def test_cover_fields_fill_only_unprotected_cover_cells(tpl, tmp_path):
    fp_docx.edit_text(tpl, "<w:r><w:t>[amount]</w:t></w:r>",
                      '<w:ins w:id="9" w:author="Officer" w:date="2026-10-01T09:00:00Z"><w:r><w:t>[amount]</w:t></w:r></w:ins>')
    r = fp_docx.build(tpl, write(tmp_path / "c.md", FIRST), tmp_path / "v1.docx", mode="first", template=tpl,
                      fields={"Borrower": "Cedar Foods", "Amount": "EUR 12m", "EUR m": "overwritten?"})
    cells = [c.text for t in docx.Document(tmp_path / "v1.docx").tables for row in t.rows for c in row.cells]
    assert "Cedar Foods" in cells                        # ordinary cover cell is filled
    assert "EUR 12m" not in cells                        # a cell with tracked changes is protected
    assert "FY2025 (template)" in cells and "overwritten?" not in cells   # tables inside sections are not cover fields
    assert sorted(r["missing_fields"]) == ["Amount", "EUR m"]


# ---------- build modes ----------

def test_first_build_needs_the_blank_template_itself(tpl, tmp_path):
    c = write(tmp_path / "c.md", FIRST)
    with pytest.raises(fp_docx.Refuse, match="--template"):
        fp_docx.build(tpl, c, tmp_path / "a.docx", mode="first")
    other = tmp_path / "other.docx"
    fp_docx.make_template(other, comment_on="market-risk")
    with pytest.raises(fp_docx.Refuse, match="identical"):
        fp_docx.build(other, c, tmp_path / "b.docx", mode="first", template=tpl)
    write(tmp_path / "template.sections.json", '{"sections": {}}')
    with pytest.raises(fp_docx.Refuse, match="assistant record"):
        fp_docx.build(tpl, c, tmp_path / "c.docx", mode="first", template=tpl)
    assert not any((tmp_path / f"{n}.docx").exists() for n in "abc")


def test_existing_needs_the_template_and_no_record_and_a_normal_round_needs_one(tpl, tmp_path):
    c = write(tmp_path / "c.md", FIRST)
    v1 = tmp_path / "v1.docx"
    fp_docx.build(tpl, c, v1, mode="first", template=tpl)
    with pytest.raises(fp_docx.Refuse, match="normal round"):
        fp_docx.build(v1, c, tmp_path / "v2.docx", mode="existing", template=tpl)
    with pytest.raises(fp_docx.Refuse, match="--existing"):
        fp_docx.build(tpl, c, tmp_path / "v3.docx")
    draft = tmp_path / "draft.docx"
    shutil.copy(tpl, draft)
    with pytest.raises(fp_docx.Refuse, match="--template"):
        fp_docx.build(draft, c, tmp_path / "v4.docx", mode="existing")


def test_existing_draft_rewrites_unprotected_sections_then_normal_rounds_follow_ownership(tpl, tmp_path):
    draft = tmp_path / "Officer-draft.docx"
    fp_docx.make_template(draft, comment_on="market-risk")
    fp_docx.edit_text(draft, '<w:r><w:t xml:space="preserve">[Summarise the proposal]</w:t></w:r>',
                      '<w:ins w:id="7" w:author="Officer" w:date="2026-10-01T09:00:00Z">'
                      "<w:r><w:t>Officer summary.</w:t></w:r></w:ins>")
    fp_docx.edit_text(draft, "[State the recommendation]", "Approve, in the officer's words.")
    before = draft.read_bytes()
    v1 = tmp_path / "FP-v01.docx"
    r = fp_docx.build(draft, write(tmp_path / "c.md", FIRST), v1, mode="existing", template=tpl,
                      protect=["recommendation"])
    assert r["written"] == ["financial-analysis"]
    assert r["frozen"] == {"summary": "has comments or tracked changes",
                           "market-risk": "has comments or tracked changes",
                           "recommendation": "protected at the officer's request"}
    md = fp_docx.to_markdown(v1)
    assert "Revenue rose 20%." in md and "FY2025 (template)" in md        # rewritten; fixed table kept
    assert "Officer summary." in md and "officer's words" in md and "[Describe market risks]" in md
    proposals = (tmp_path / "FP-v01.proposals.md").read_text()
    assert "Approve." in proposals and "Milk prices rose." in proposals
    side = json.loads((tmp_path / "FP-v01.sections.json").read_text())["sections"]
    assert {k: v["owner"] for k, v in side.items()} == {
        "summary": "other", "financial-analysis": "ai", "market-risk": "other", "recommendation": "other"}
    assert draft.read_bytes() == before                                    # the officer's file is untouched
    r2 = fp_docx.build(v1, write(tmp_path / "c2.md", "## financial-analysis\nRevenue rose 25%.\n\n{{keep:1}}\n\n"
                                 "## recommendation\nDecline.\n"), tmp_path / "FP-v02.docx")
    assert r2["written"] == ["financial-analysis"]
    assert r2["frozen"] == {"recommendation": "not written by the assistant in the base version"}
    assert "Decline." in (tmp_path / "FP-v02.proposals.md").read_text()


# ---------- edit detection through a real Word save ----------

WORD = Path(__file__).parent / "fixtures" / "word-saved"   # synthetic FP built by the tool, then saved by Word


@pytest.mark.parametrize("version", [1, 2])
def test_real_word_save_changes_only_the_edited_section(version):
    before = fp_docx.signatures(fp_docx.Doc(WORD / f"assistant-v{version}.docx"))
    after = fp_docx.signatures(fp_docx.Doc(WORD / f"word-saved-v{version}.docx"))
    assert list(before) == list(after) == ["summary", "financial-analysis", "market-risk", "recommendation"]
    assert [sid for sid in before if before[sid] != after[sid]] == ["recommendation"]


def test_rounds_after_a_real_word_save_write_untouched_sections_and_keep_the_officers_text(tmp_path):
    base, v2, v3 = (tmp_path / f"FP-v0{n}.docx" for n in (1, 2, 3))
    shutil.copy(WORD / "word-saved-v1.docx", base)
    record = {sid: {"owner": "ai", "signature": sig, "ids": []}
              for sid, sig in fp_docx.signatures(fp_docx.Doc(WORD / "assistant-v1.docx")).items()}
    write(tmp_path / "FP-v01.sections.json", json.dumps({"base": "template.docx", "sections": record}))
    r = fp_docx.build(base, write(tmp_path / "c.md", "## summary\nNew summary.\n\n## recommendation\nDecline.\n"),
                      v2, template=WORD / "template.docx", fields={"Amount": "EUR 15m"})
    assert r["written"] == ["summary"]
    assert r["frozen"] == {"recommendation": "edited since the assistant wrote it"}
    md = fp_docx.to_markdown(v2)
    assert "New summary." in md and "wow" in md and "Decline." not in md
    assert "EUR 15m" in [c.text for t in docx.Document(v2).tables for row in t.rows for c in row.cells]
    # The assistant tries the officer's section again in the next round: still frozen, "wow" still there.
    r = fp_docx.build(v2, write(tmp_path / "c3.md", "## recommendation\nDecline.\n\n## market-risk\nPrices fell.\n"),
                      v3, template=WORD / "template.docx")
    assert r["written"] == ["market-risk"]
    assert r["frozen"] == {"recommendation": "not written by the assistant in the base version"}
    md = fp_docx.to_markdown(v3)
    assert "wow" in md and "Prices fell." in md and "Decline." not in md


APPROVE = '<w:r><w:t xml:space="preserve">Approve.</w:t></w:r>'
FORMAT_EDITS = {
    "highlight": APPROVE.replace("<w:r>", '<w:r><w:rPr><w:highlight w:val="yellow"/></w:rPr>'),
    "colour": APPROVE.replace("<w:r>", '<w:r><w:rPr><w:color w:val="FF0000"/></w:rPr>'),
    "strike": APPROVE.replace("<w:r>", "<w:r><w:rPr><w:strike/></w:rPr>"),
    "centre": '<w:pPr><w:jc w:val="center"/></w:pPr>' + APPROVE,
    "hyperlink": f'<w:hyperlink w:anchor="annex">{APPROVE}</w:hyperlink>',
    "footnote": APPROVE + '<w:r><w:footnoteReference w:id="1"/></w:r>',
    "break": APPROVE.replace("Appr", 'Appr</w:t><w:br/><w:t xml:space="preserve">'),
    "tab": APPROVE.replace("Appr", 'Appr</w:t><w:tab/><w:t xml:space="preserve">'),
}


@pytest.mark.parametrize("edited", FORMAT_EDITS.values(), ids=FORMAT_EDITS.keys())
def test_formatting_edit_freezes_the_section(tpl, tmp_path, edited):
    v1 = tmp_path / "v1.docx"
    fp_docx.build(tpl, write(tmp_path / "c.md", FIRST), v1, mode="first", template=tpl)
    fp_docx.edit_text(v1, APPROVE, edited)
    r = fp_docx.build(v1, write(tmp_path / "c2.md", "## recommendation\nDecline.\n"), tmp_path / "v2.docx")
    assert r["frozen"] == {"recommendation": "edited since the assistant wrote it"}


def test_generated_paragraphs_have_no_explicit_default_style(tpl, tmp_path):
    out = tmp_path / "v1.docx"
    fp_docx.build(tpl, write(tmp_path / "c.md", FIRST), out, mode="first", template=tpl)
    with zipfile.ZipFile(out) as z:
        xml = z.read("word/document.xml").decode()
    assert "<w:p><w:r><w:t xml:space=\"preserve\">Approve.</w:t></w:r></w:p>" in xml
    assert 'w:val="Normal"' not in xml


# ---------- preservation ----------

def test_page_breaks_notes_and_bookmarks_in_a_written_section_survive(tpl, tmp_path):
    guidance = '<w:t xml:space="preserve">[Summarise the proposal]</w:t></w:r></w:p>'
    fp_docx.edit_text(tpl, guidance, guidance + '<w:p><w:r><w:br w:type="page"/></w:r></w:p>'
                      '<w:p><w:r><w:t>Source</w:t></w:r><w:r><w:footnoteReference w:id="1"/></w:r></w:p>'
                      '<w:p><w:r><w:t>Anchor</w:t></w:r><w:bookmarkEnd w:id="3"/></w:p>')
    fp_docx.edit_text(tpl, "<w:body>", '<w:body><w:p><w:bookmarkStart w:id="3" w:name="Cover"/></w:p>')
    out = tmp_path / "v1.docx"
    r = fp_docx.build(tpl, write(tmp_path / "c.md", FIRST), out, mode="first", template=tpl)
    assert "summary" in r["written"]
    with zipfile.ZipFile(out) as z:
        xml = z.read("word/document.xml").decode()
    assert "[Summarise the proposal]" not in xml and "Cedar Foods seeks EUR 12m." in xml
    assert '<w:br w:type="page"/>' in xml and '<w:footnoteReference w:id="1"/>' in xml
    assert '<w:bookmarkEnd w:id="3"/>' in xml


def test_bookmark_in_a_cover_value_cell_protects_it(tpl, tmp_path):
    fp_docx.edit_text(tpl, "<w:r><w:t>[name]</w:t></w:r>",
                      '<w:bookmarkStart w:id="4" w:name="BorrowerName"/><w:r><w:t>[name]</w:t></w:r>'
                      '<w:bookmarkEnd w:id="4"/>')
    out = tmp_path / "v1.docx"
    r = fp_docx.build(tpl, write(tmp_path / "c.md", FIRST), out, mode="first", template=tpl,
                      fields={"Borrower": "Cedar Foods", "Amount": "EUR 12m"})
    assert r["protected_fields"] == ["Borrower"] and r["missing_fields"] == ["Borrower"]
    cells = [c.text for t in docx.Document(out).tables for row in t.rows for c in row.cells]
    assert "[name]" in cells and "EUR 12m" in cells


def test_build_that_loses_a_kept_object_is_refused_and_writes_nothing(tmp_path, monkeypatch):
    tpl = tmp_path / "template.docx"
    fp_docx.make_template(tpl, comment_on="market-risk")      # a frozen section, so proposals would be due
    monkeypatch.setattr(fp_docx, "render", lambda doc, sec, items, inc: fp_docx.Gen(doc).p("Lost the table.", None))
    with pytest.raises(fp_docx.Refuse, match="nothing written"):
        fp_docx.build(tpl, write(tmp_path / "c.md", FIRST), tmp_path / "FP-v01.docx", mode="first", template=tpl)
    assert sorted(p.name for p in tmp_path.iterdir()) == ["c.md", "template.docx"]


def test_build_that_changes_a_section_it_did_not_write_is_refused(tmp_path, monkeypatch):
    tpl = tmp_path / "template.docx"
    fp_docx.make_template(tpl, comment_on="market-risk")
    real = fp_docx.splice
    monkeypatch.setattr(fp_docx, "splice", lambda xml, edits: real(xml, edits).replace(b"[Describe", b"[Rewrite"))
    with pytest.raises(fp_docx.Refuse, match="market-risk"):
        fp_docx.build(tpl, write(tmp_path / "c.md", FIRST), tmp_path / "FP-v01.docx", mode="first", template=tpl)
    assert sorted(p.name for p in tmp_path.iterdir()) == ["c.md", "template.docx"]


# ---------- styles for new text ----------

def edit_part(path, name, old, new):
    with zipfile.ZipFile(path) as z:
        parts = {i.filename: z.read(i.filename) for i in z.infolist()}
    assert old.encode() in parts[name], old
    parts[name] = parts[name].replace(old.encode(), new.encode(), 1)
    with zipfile.ZipFile(path, "w") as z:
        for n, data in parts.items():
            z.writestr(n, data)


def paragraph_xml(path, text):
    with zipfile.ZipFile(path) as z:
        xml = z.read("word/document.xml").decode()
    return next(p for p in re.findall(r"<w:p>.*?</w:p>|<w:p .*?</w:p>", xml) if text in p)


def test_prose_never_takes_a_list_style_from_the_section(tpl, tmp_path):
    edit_part(tpl, "word/styles.xml", "</w:styles>", '<w:style w:type="paragraph" w:styleId="Points">'
              '<w:name w:val="Points"/><w:pPr><w:numPr><w:numId w:val="3"/></w:numPr></w:pPr></w:style></w:styles>')
    fp_docx.edit_text(tpl, '<w:pStyle w:val="Normal"/></w:pPr><w:r><w:t xml:space="preserve">[Describe',
                      '<w:pStyle w:val="ListBullet"/></w:pPr><w:r><w:t xml:space="preserve">[Describe')
    fp_docx.edit_text(tpl, '<w:pStyle w:val="Normal"/></w:pPr><w:r><w:t xml:space="preserve">[Summarise',
                      '<w:pStyle w:val="Points"/></w:pPr><w:r><w:t xml:space="preserve">[Summarise')
    out = tmp_path / "v1.docx"
    fp_docx.build(tpl, write(tmp_path / "c.md", "## summary\nCedar Foods seeks EUR 12m.\n\n"
                             "## market-risk\nPrices rose.\n\n- Milk prices rose.\n"), out, mode="first", template=tpl)
    assert "pStyle" not in paragraph_xml(out, "Cedar Foods seeks")      # default style, not Points
    assert "pStyle" not in paragraph_xml(out, "Prices rose.")           # default style, not List Bullet
    assert 'w:val="ListBullet"' in paragraph_xml(out, "Milk prices rose.")


def test_tables_next_to_a_generated_table_are_kept_apart(tpl, tmp_path):
    out = tmp_path / "v1.docx"
    fp_docx.build(tpl, write(tmp_path / "c.md", "## financial-analysis\n| A | B |\n|---|---|\n| 1 | 2 |\n\n"
                             "{{keep:1}}\n\n| C | D |\n|---|---|\n| 3 | 4 |\n"), out, mode="first", template=tpl)
    with zipfile.ZipFile(out) as z:
        xml = z.read("word/document.xml").decode()
    assert "</w:tbl><w:tbl>" not in xml and xml.count("</w:tbl><w:p/><w:tbl>") == 2
    assert len(docx.Document(out).tables) == 4


# ---------- reporting ----------

@pytest.mark.parametrize("heading", ["3. Recommendation", "3.Recommendation", "3 Recommendation",
                                      "3. Recommendation and conditions"])
def test_missing_template_sections_are_reported_and_text_under_them_is_kept(tpl, tmp_path, heading):
    draft = tmp_path / "draft.docx"   # the officer typed a bold paragraph, not a heading
    shutil.copy(tpl, draft)
    fp_docx.edit_text(draft, '<w:pStyle w:val="Heading1"/></w:pPr><w:r><w:t xml:space="preserve">3. Recommendation',
                      '<w:pStyle w:val="Normal"/></w:pPr><w:r><w:rPr><w:b/></w:rPr>'
                      f'<w:t xml:space="preserve">{heading}')
    fp_docx.edit_text(draft, "[State the recommendation]", "Approve with two covenants (officer text).")
    out = tmp_path / "v1.docx"
    r = fp_docx.build(draft, write(tmp_path / "c.md", FIRST), out, mode="existing", template=tpl)
    assert r["missing_sections"] == ["recommendation"] and r["unmatched"] == ["recommendation"]
    assert "recommendation" in r["frozen"]["market-risk"] and r["complete"] is False
    md = fp_docx.to_markdown(out)
    assert heading in md and "Approve with two covenants (officer text)." in md
    proposals = (tmp_path / "v1.proposals.md").read_text()
    assert "Milk prices rose." in proposals and "Approve." in proposals
    assert fp_docx.check(out, tpl)["missing_sections"] == ["recommendation"]
    ok = fp_docx.build(tpl, write(tmp_path / "c2.md", FIRST), tmp_path / "v2.docx", mode="first", template=tpl)
    assert ok["missing_sections"] == [] and ok["unmatched"] == [] and ok["frozen"] == {}
    assert ok["complete"] is True


# ---------- robustness ----------

def test_draft_with_a_byte_order_mark_builds(tpl, tmp_path):
    c = tmp_path / "c.md"
    c.write_bytes(b"\xef\xbb\xbf" + FIRST.encode("utf-8"))
    r = fp_docx.build(tpl, c, tmp_path / "v1.docx", mode="first", template=tpl)
    assert r["written"] == ["summary", "financial-analysis", "market-risk", "recommendation"]
    assert r["warnings"] == []


def test_duplicate_draft_sections_are_refused_and_stray_text_is_reported(tpl, tmp_path):
    with pytest.raises(fp_docx.Refuse, match="summary"):
        fp_docx.build(tpl, write(tmp_path / "c.md", "## summary\nA.\n\n## Summary\nB.\n"), tmp_path / "v1.docx",
                      mode="first", template=tpl)
    assert not (tmp_path / "v1.docx").exists()
    r = fp_docx.build(tpl, write(tmp_path / "c2.md", "Draft for Cedar Foods\n\n## summary\nA.\n"), tmp_path / "v2.docx",
                      mode="first", template=tpl)
    assert r["written"] == ["summary"] and "Draft for Cedar Foods" in r["warnings"][0]
    assert r["complete"] is False


def test_cli_refuses_bad_input_with_exit_3_and_no_traceback(tpl, tmp_path, capsys):
    c = write(tmp_path / "c.md", FIRST)
    first = ["build", "--base", str(tpl), "--first", "--template", str(tpl)]
    (tmp_path / "bad.json").write_bytes(b'{"Borrower": ')
    (tmp_path / "list.json").write_bytes(b'["Borrower"]')
    (tmp_path / "latin1.md").write_bytes("## summary\nCaf\xe9 au lait.\n".encode("latin-1"))
    assert fp_docx.main(first + ["--content", str(c), "--out", str(tmp_path / "a.docx"),
                                 "--fields", str(tmp_path / "bad.json")]) == 3
    assert fp_docx.main(first + ["--content", str(c), "--out", str(tmp_path / "b.docx"),
                                 "--fields", str(tmp_path / "list.json")]) == 3
    assert fp_docx.main(first + ["--content", str(tmp_path / "latin1.md"), "--out", str(tmp_path / "c.docx")]) == 3
    assert fp_docx.main(first + ["--content", str(c), "--out", str(tmp_path / "v1.docx")]) == 0
    write(tmp_path / "v1.sections.json", "{not json")
    assert fp_docx.main(["build", "--base", str(tmp_path / "v1.docx"), "--content", str(c),
                         "--out", str(tmp_path / "v2.docx")]) == 3
    err = capsys.readouterr().err
    assert err.count("REFUSED: ") == 4 and "Traceback" not in err
    assert not any((tmp_path / f"{n}.docx").exists() for n in ("a", "b", "c", "v2"))


def test_outputs_are_never_overwritten_and_are_readable_by_others(tpl, tmp_path):
    c = write(tmp_path / "c.md", FIRST + "\n## annex-9\nText\n")
    write(tmp_path / "v1.sections.json", "{}")
    with pytest.raises(fp_docx.Refuse, match="v1.sections.json"):
        fp_docx.build(tpl, c, tmp_path / "v1.docx", mode="first", template=tpl)
    write(tmp_path / "v2.proposals.md", "officer notes")
    with pytest.raises(fp_docx.Refuse, match="v2.proposals.md"):
        fp_docx.build(tpl, c, tmp_path / "v2.docx", mode="first", template=tpl)
    assert (tmp_path / "v1.sections.json").read_text() == "{}"
    assert (tmp_path / "v2.proposals.md").read_text() == "officer notes"
    assert not (tmp_path / "v1.docx").exists() and not (tmp_path / "v2.docx").exists()
    fp_docx.build(tpl, c, tmp_path / "v3.docx", mode="first", template=tpl)
    for name in ("v3.docx", "v3.sections.json", "v3.proposals.md"):
        assert (tmp_path / name).stat().st_mode & 0o777 == 0o644, name
    assert not list(tmp_path.glob("*.tmp"))


def test_tabs_and_line_breaks_are_text_and_numbered_headings_keep_their_id(tpl):
    fp_docx.edit_text(tpl, '<w:pStyle w:val="Heading1"/></w:pPr><w:r><w:t xml:space="preserve">1. Summary</w:t>',
                      '<w:pStyle w:val="Heading1"/><w:tabs><w:tab w:val="left" w:pos="720"/></w:tabs></w:pPr>'
                      "<w:r><w:t>1.</w:t><w:tab/><w:t>Summary</w:t>")
    fp_docx.edit_text(tpl, '<w:t xml:space="preserve">[Explain the figures]</w:t>',
                      '<w:t>Line one</w:t><w:br/><w:t>Line</w:t><w:tab/><w:t>two</w:t><w:br w:type="page"/>')
    doc = fp_docx.Doc(tpl)
    assert doc.sections[0]["id"] == "summary" and doc.sections[0]["heading"] == "1.\tSummary"
    assert doc.sections[1]["body"][0].plain() == "Line one\nLine\ttwo"


def test_markdown_shows_comments_in_a_read_only_note(tmp_path):
    tpl = tmp_path / "template.docx"
    fp_docx.make_template(tpl, comment_on="market-risk")
    edit_part(tpl, "word/comments.xml", "Check the FX exposure", "Check the FX exposure --> and hedging")
    v1 = tmp_path / "v1.docx"
    fp_docx.build(tpl, write(tmp_path / "c.md", FIRST), v1, mode="first", template=tpl)
    md = fp_docx.to_markdown(v1)
    assert "<!-- comments" in md and "Check the FX exposure" in md
    r = fp_docx.build(v1, write(tmp_path / "again.md", md), tmp_path / "v2.docx", template=tpl)
    assert r["unmatched"] == [] and r["warnings"] == []
    assert "hedging" not in "\n".join(p.text for p in docx.Document(tmp_path / "v2.docx").paragraphs)


def test_dotx_template_builds_a_document_and_macro_files_are_refused(tmp_path):
    dotx = tmp_path / "template.dotx"
    fp_docx.make_template(dotx)
    edit_part(dotx, "[Content_Types].xml", "wordprocessingml.document.main+xml", "wordprocessingml.template.main+xml")
    out = tmp_path / "v1.docx"
    fp_docx.build(dotx, write(tmp_path / "c.md", FIRST), out, mode="first", template=dotx)
    with zipfile.ZipFile(out) as z:
        types = z.read("[Content_Types].xml").decode()
    assert "wordprocessingml.document.main+xml" in types and "template.main" not in types
    assert docx.Document(out).paragraphs                     # python-docx refuses a template content type
    with pytest.raises(fp_docx.Refuse, match="--first"):
        fp_docx.build(dotx, tmp_path / "c.md", tmp_path / "v2.docx", mode="existing")
    docm = tmp_path / "draft.docm"
    fp_docx.make_template(docm)
    edit_part(docm, "[Content_Types].xml", "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
              "application/vnd.ms-word.document.macroEnabled")
    with pytest.raises(fp_docx.Refuse, match="macro"):
        fp_docx.build(docm, tmp_path / "c.md", tmp_path / "v3.docx", mode="existing")


def test_edit_text_raises_value_error_when_the_text_is_absent(tpl):
    with pytest.raises(ValueError):
        fp_docx.edit_text(tpl, "not in the document", "x")


# ---------- fixes after the Task 2a review ----------

def unstyle_headings(path, *texts):
    for text in texts:
        for style in ("Heading1", "Heading2"):
            try:
                fp_docx.edit_text(path, f'<w:pStyle w:val="{style}"/></w:pPr><w:r><w:t xml:space="preserve">{text}',
                                  f'<w:pStyle w:val="Normal"/></w:pPr><w:r><w:t xml:space="preserve">{text}')
                break
            except ValueError:
                continue


def test_fields_are_refused_when_the_cover_cannot_be_told_from_the_body(tpl, tmp_path):
    c = write(tmp_path / "c.md", "## market-risk\nM.\n")
    flat = tmp_path / "flat.docx"                     # no recognised headings at all
    shutil.copy(tpl, flat)
    unstyle_headings(flat, "1. Summary", "2. Financial analysis", "2.1 Market risk", "3. Recommendation")
    with pytest.raises(fp_docx.Refuse, match="no recognised headings"):
        fp_docx.build(flat, c, tmp_path / "a.docx", mode="existing", template=tpl, fields={"Amount": "EUR 15m"})
    late = tmp_path / "late.docx"                     # the first two headings are not recognised
    shutil.copy(tpl, late)
    unstyle_headings(late, "1. Summary", "2. Financial analysis")
    with pytest.raises(fp_docx.Refuse, match="market-risk.*summary"):
        fp_docx.build(late, c, tmp_path / "b.docx", mode="existing", template=tpl, fields={"EUR m": "x"})
    assert not (tmp_path / "a.docx").exists() and not (tmp_path / "b.docx").exists()
    r = fp_docx.build(late, c, tmp_path / "c.docx", mode="existing", template=tpl)   # without --fields it builds
    assert r["missing_sections"] == ["summary", "financial-analysis"]


def test_complete_is_false_unless_everything_was_checked_and_applied(tpl, tmp_path):
    v1 = tmp_path / "v1.docx"
    r = fp_docx.build(tpl, write(tmp_path / "c.md", FIRST), v1, mode="first", template=tpl,
                      fields={"Borrower": "Cedar Foods", "Tenor": "5 years"})
    assert r["frozen"] == {} and r["unmatched"] == [] and r["missing_sections"] == []
    assert r["missing_fields"] == ["Tenor"] and r["complete"] is False
    c2 = write(tmp_path / "c2.md", "## summary\nS.\n")
    r = fp_docx.build(v1, c2, tmp_path / "v2.docx")
    assert r["missing_sections"] is None and r["complete"] is False
    assert any("not checked" in w for w in r["warnings"])
    assert fp_docx.check(tmp_path / "v2.docx")["missing_sections"] is None
    r = fp_docx.build(v1, c2, tmp_path / "v3.docx", template=tpl)
    assert r["missing_sections"] == [] and r["warnings"] == [] and r["complete"] is True


def template_table(text):
    return ('<w:tbl><w:tblPr><w:tblW w:w="0" w:type="auto"/></w:tblPr><w:tblGrid><w:gridCol w:w="9000"/>'
            f"</w:tblGrid><w:tr><w:tc><w:p><w:r><w:t>{text}</w:t></w:r></w:p></w:tc></w:tr></w:tbl>")


def test_kept_tables_never_end_up_touching_unless_they_touched_in_the_base(tpl, tmp_path):
    fin_end = "<w:t>FY2025 (template)</w:t></w:r></w:p></w:tc></w:tr></w:tbl>"
    fp_docx.edit_text(tpl, fin_end, fin_end + '<w:p><w:r><w:t>[Comment on the balance sheet]</w:t></w:r></w:p>'
                      + template_table("Balance sheet (template)"))
    guidance = '<w:t xml:space="preserve">[Summarise the proposal]</w:t></w:r></w:p>'
    fp_docx.edit_text(tpl, guidance, guidance + template_table("Fees A") + template_table("Fees B"))
    out = tmp_path / "v1.docx"
    fp_docx.build(tpl, write(tmp_path / "c.md", "## summary\nS.\n\n## financial-analysis\nRevenue rose 20%.\n\n"
                             "The balance sheet is strong.\n"), out, mode="first", template=tpl)
    with zipfile.ZipFile(out) as z:
        xml = z.read("word/document.xml").decode()
    fin = xml[xml.index("2. Financial analysis"):xml.index("2.1 Market risk")]
    assert "</w:tbl><w:tbl>" not in fin and "</w:tbl><w:p/><w:tbl>" in fin
    summary = xml[xml.index("1. Summary"):xml.index("2. Financial analysis")]
    assert "Fees A" in summary and "</w:tbl><w:tbl>" in summary     # they already touched in the template


def test_protected_section_belongs_to_the_officer_afterwards(tpl, tmp_path):
    v1, v2, v3 = (tmp_path / f"v{n}.docx" for n in (1, 2, 3))
    fp_docx.build(tpl, write(tmp_path / "c.md", FIRST), v1, mode="first", template=tpl)
    fp_docx.build(v1, write(tmp_path / "c2.md", "## summary\nNew summary.\n"), v2, protect=["market-risk"])
    side = json.loads((tmp_path / "v2.sections.json").read_text())["sections"]
    assert side["market-risk"]["owner"] == "other" and side["recommendation"]["owner"] == "ai"
    r = fp_docx.build(v2, write(tmp_path / "c3.md", "## market-risk\nM.\n"), v3)
    assert r["frozen"] == {"market-risk": "not written by the assistant in the base version"}


def test_cover_field_fill_that_loses_its_formatting_is_refused(tpl, tmp_path, monkeypatch):
    fp_docx.edit_text(tpl, "<w:r><w:t>[name]</w:t></w:r>", "<w:r><w:rPr><w:b/></w:rPr><w:t>[name]</w:t></w:r>")
    real = fp_docx.fill_fields

    def lossy(doc, xml, fields, inc):
        new, *rest = real(doc, xml, fields, inc)
        return (new.replace(b"<w:rPr><w:b/></w:rPr>", b"", 1), *rest)

    monkeypatch.setattr(fp_docx, "fill_fields", lossy)
    with pytest.raises(fp_docx.Refuse, match="nothing written"):
        fp_docx.build(tpl, write(tmp_path / "c.md", FIRST), tmp_path / "v1.docx", mode="first", template=tpl,
                      fields={"Borrower": "Cedar Foods"})
    assert not (tmp_path / "v1.docx").exists() and not (tmp_path / "v1.sections.json").exists()


# ---------- Task 2b: table formatting ----------

def edit_xml(path, pattern, repl):
    """Like edit_text, with a regular expression (first match only)."""
    with zipfile.ZipFile(path) as z:
        parts = {i.filename: z.read(i.filename) for i in z.infolist()}
    new, n = re.subn(pattern.encode(), repl.encode(), parts["word/document.xml"], count=1)
    assert n == 1, pattern
    parts["word/document.xml"] = new
    with zipfile.ZipFile(path, "w") as z:
        for name, data in parts.items():
            z.writestr(name, data)


GEN_ROW = "</w:tblGrid><w:tr><w:tc><w:tcPr>"           # first row of the assistant table only
GEN_PR = '<w:tblW w:w="0" w:type="auto"/><w:tblLook'
WORD_LOOK = ('<w:tblLook w:val="04A0" w:firstRow="1" w:lastRow="0" w:firstColumn="1" '
             'w:lastColumn="0" w:noHBand="0" w:noVBand="1"/>')
TABLE_EDITS = {  # name: (pattern, replacement, freezes)
    "shading": ("</w:tcPr>", '<w:shd w:val="clear" w:color="auto" w:fill="FFFF00"/></w:tcPr>', True),
    "table style": ('<w:tblStyle w:val="TableGrid"/>', '<w:tblStyle w:val="LightShading"/>', True),
    "header row": (GEN_ROW, "</w:tblGrid><w:tr><w:trPr><w:tblHeader/></w:trPr><w:tc><w:tcPr>", True),
    "column width": ('<w:tcW w:w="[0-9]+" w:type="[a-z]+"/>', '<w:tcW w:w="2400" w:type="dxa"/>', True),
    "grid": ('<w:gridCol w:w="4500"/><w:gridCol w:w="4500"/>' + GEN_ROW,
             '<w:gridCol w:w="1031"/><w:gridCol w:w="1343"/>' + GEN_ROW, False),
    "indent": (GEN_PR, '<w:tblW w:w="0" w:type="auto"/><w:tblInd w:w="10" w:type="dxa"/><w:tblLook', False),
    "cell margins": (GEN_PR, '<w:tblW w:w="0" w:type="auto"/><w:tblCellMar><w:left w:w="10" w:type="dxa"/>'
                     '</w:tblCellMar><w:tblLook', False),
    "zero look": (WORD_LOOK, '<w:tblLook w:val="0000" w:firstRow="0" w:noVBand="0"/>', True),
    "header look off": (WORD_LOOK, WORD_LOOK.replace('w:firstRow="1"', 'w:firstRow="0"'), True),
    "look hex only": (WORD_LOOK, '<w:tblLook w:val="04A0"/>', False),
    "nil width": (GEN_PR, '<w:tblW w:w="0" w:type="nil"/><w:tblLook', False),
    "row exceptions": (GEN_ROW, "</w:tblGrid><w:tr><w:tblPrEx><w:tblCellMar><w:top w:w=\"0\" w:type=\"dxa\"/>"
                       "</w:tblCellMar></w:tblPrEx><w:tc><w:tcPr>", False),
}


@pytest.mark.parametrize("pattern,repl,freezes", TABLE_EDITS.values(), ids=TABLE_EDITS.keys())
def test_table_formatting_edits_freeze_but_word_layout_noise_does_not(tpl, tmp_path, pattern, repl, freezes):
    v1 = tmp_path / "v1.docx"
    fp_docx.build(tpl, write(tmp_path / "c.md", FIRST), v1, mode="first", template=tpl)
    edit_xml(v1, pattern, repl)
    r = fp_docx.build(v1, write(tmp_path / "c2.md", "## financial-analysis\nRevenue rose 25%.\n\n{{keep:1}}\n"),
                      tmp_path / "v2.docx")
    assert r["frozen"] == ({"financial-analysis": "edited since the assistant wrote it"} if freezes else {})


def test_generated_tables_carry_no_explicit_widths(tpl, tmp_path):
    out = tmp_path / "v1.docx"
    fp_docx.build(tpl, write(tmp_path / "c.md", FIRST), out, mode="first", template=tpl)
    with zipfile.ZipFile(out) as z:
        xml = z.read("word/document.xml").decode()
    table = xml[xml.index("<w:tblStyle"):xml.index("</w:tbl>", xml.index("<w:tblStyle"))]
    assert 'w:type="dxa"' not in table and '<w:tcW w:w="0" w:type="auto"/>' in table
    assert WORD_LOOK + '<w:tblDescription' in table


# ---------- Task 2b: hidden bookmarks ----------

HLK = "_Hlk180000001"


def add_part(path, name, data):
    with zipfile.ZipFile(path, "a") as z:
        z.writestr(name, data)


def officer_draft_with_hidden_bookmark(tpl, path):
    shutil.copy(tpl, path)
    fp_docx.edit_text(path, '<w:r><w:t xml:space="preserve">[Summarise the proposal]</w:t></w:r>',
                      f'<w:bookmarkStart w:id="5" w:name="{HLK}"/><w:r><w:t>Officer summary.</w:t></w:r>'
                      '<w:bookmarkEnd w:id="5"/>')


def test_paragraph_with_an_unreferenced_hidden_bookmark_is_replaced(tpl, tmp_path):
    draft, out = tmp_path / "draft.docx", tmp_path / "v1.docx"
    officer_draft_with_hidden_bookmark(tpl, draft)
    r = fp_docx.build(draft, write(tmp_path / "c.md", "## summary\nNew summary.\n"), out,
                      mode="existing", template=tpl)
    assert r["written"] == ["summary"]
    md = fp_docx.to_markdown(out)
    assert "New summary." in md and "Officer summary." not in md      # replaced, not duplicated


FIELD = ('<w:p><w:r><w:fldChar w:fldCharType="begin"/></w:r><w:r><w:instrText xml:space="preserve"> HYPERLINK \\l "'
         + HLK + '" </w:instrText></w:r><w:r><w:fldChar w:fldCharType="separate"/></w:r><w:r><w:t>see the summary'
         '</w:t></w:r><w:r><w:fldChar w:fldCharType="end"/></w:r></w:p>')
ANCHOR = f'<w:p><w:hyperlink w:anchor="{HLK}"><w:r><w:t>see the summary</w:t></w:r></w:hyperlink></w:p>'
NOTES = (f'<?xml version="1.0" encoding="UTF-8"?><w:footnotes xmlns:w="{fp_docx.W}"><w:footnote w:id="1"><w:p>'
         f'<w:fldSimple w:instr=" PAGEREF {HLK} \\h "><w:r><w:t>1</w:t></w:r></w:fldSimple></w:p></w:footnote>'
         "</w:footnotes>")


@pytest.mark.parametrize("where", ["field", "anchor", "footnote"])
def test_paragraph_with_a_referenced_hidden_bookmark_stays(tpl, tmp_path, where):
    draft, out = tmp_path / "draft.docx", tmp_path / "v1.docx"
    officer_draft_with_hidden_bookmark(tpl, draft)
    end = "[State the recommendation]</w:t></w:r></w:p>"
    if where == "footnote":
        add_part(draft, "word/footnotes.xml", NOTES)
    else:
        fp_docx.edit_text(draft, end, end + (FIELD if where == "field" else ANCHOR))
    fp_docx.build(draft, write(tmp_path / "c.md", "## summary\nNew summary.\n"), out, mode="existing", template=tpl)
    md = fp_docx.to_markdown(out)
    assert "New summary." in md and "Officer summary." in md          # the bookmark target is kept


def test_hidden_bookmark_markers_are_not_edits_but_text_changes_are(tpl, tmp_path):
    v1 = tmp_path / "v1.docx"
    fp_docx.build(tpl, write(tmp_path / "c.md", FIRST), v1, mode="first", template=tpl)
    signature = lambda: fp_docx.signatures(fp_docx.Doc(v1))["recommendation"]  # noqa: E731
    before = signature()
    marked = f'<w:bookmarkStart w:id="8" w:name="{HLK}"/>{APPROVE}<w:bookmarkEnd w:id="8"/>'
    fp_docx.edit_text(v1, APPROVE, marked)
    assert signature() == before
    fp_docx.edit_text(v1, marked, APPROVE)
    assert signature() == before
    fp_docx.edit_text(v1, APPROVE, marked.replace("Approve.", "Approve now."))
    assert signature() != before


def test_cover_cell_with_an_unreferenced_hidden_bookmark_can_be_filled(tpl, tmp_path):
    fp_docx.edit_text(tpl, "<w:r><w:t>[name]</w:t></w:r>",
                      f'<w:bookmarkStart w:id="6" w:name="{HLK}"/><w:r><w:t>[name]</w:t></w:r><w:bookmarkEnd w:id="6"/>')
    r = fp_docx.build(tpl, write(tmp_path / "c.md", FIRST), tmp_path / "v1.docx", mode="first", template=tpl,
                      fields={"Borrower": "Cedar Foods"})
    assert r["missing_fields"] == [] and r["protected_fields"] == []


# ---------- Task 2b: Save As recovery ----------

def test_save_as_copy_is_matched_to_its_record_and_keeps_the_officers_edits(tpl, tmp_path, capsys):
    v1, v2 = tmp_path / "FP-v01.docx", tmp_path / "FP-v02.docx"
    fp_docx.build(tpl, write(tmp_path / "c1.md", FIRST), v1, mode="first", template=tpl)
    fp_docx.build(v1, write(tmp_path / "c2.md", "## summary\nCedar Foods seeks EUR 15m.\n"), v2, template=tpl)
    (tmp_path / "officer").mkdir()
    copy = tmp_path / "officer" / "FP final.docx"                   # Save As: no record next to it
    shutil.copy(v2, copy)
    fp_docx.edit_text(copy, "Revenue rose 20%.", "Revenue rose 20% (officer).")
    fp_docx.edit_text(copy, "Approve.", "Approve, subject to covenants.")
    found = fp_docx.find_record(copy, folder=tmp_path)
    assert found["source"] == str(tmp_path / "FP-v02.sections.json") and found["ambiguous"] is False
    assert found["candidates"][0] == {"record": found["source"], "matching": 2, "total": 4, "same_ids": True}
    assert fp_docx.main(["find-record", str(copy), "--folder", str(tmp_path)]) == 0
    assert json.loads(capsys.readouterr().out)["source"] == found["source"]
    c3 = write(tmp_path / "c3.md", "## summary\nS.\n\n## financial-analysis\nF.\n\n{{keep:1}}\n\n## market-risk\nM.\n\n"
                                   "## recommendation\nR.\n")
    with pytest.raises(fp_docx.Refuse, match="find-record.*--record.*--existing --template"):
        fp_docx.build(copy, c3, tmp_path / "FP-v03.docx", template=tpl)
    with pytest.raises(fp_docx.Refuse, match="--record"):
        fp_docx.build(copy, c3, tmp_path / "FP-v03.docx", mode="existing", template=tpl, record=found["source"])
    r = fp_docx.build(copy, c3, tmp_path / "FP-v03.docx", template=tpl, record=found["source"])
    assert r["written"] == ["summary", "market-risk"]
    assert r["frozen"] == {"financial-analysis": "edited since the assistant wrote it",
                           "recommendation": "edited since the assistant wrote it"}
    md = fp_docx.to_markdown(tmp_path / "FP-v03.docx")
    assert "20% (officer)" in md and "subject to covenants" in md


def test_find_record_is_ambiguous_when_two_records_fit_equally(tpl, tmp_path):
    v1 = tmp_path / "v1.docx"
    fp_docx.build(tpl, write(tmp_path / "c.md", FIRST), v1, mode="first", template=tpl)
    shutil.copy(tmp_path / "v1.sections.json", tmp_path / "v1 backup.sections.json")
    copy = tmp_path / "renamed.docx"
    shutil.copy(v1, copy)
    found = fp_docx.find_record(copy)
    assert found["ambiguous"] is True and found["source"] is None and len(found["candidates"]) == 2


def test_first_draft_without_records_is_ambiguous_and_can_explicitly_use_existing(tpl, tmp_path):
    deal = tmp_path / "deal"
    deal.mkdir()
    draft = deal / "Officer draft.docx"
    shutil.copy(tpl, draft)
    fp_docx.edit_text(draft, "[State the recommendation]", "Approve (officer).")
    assert fp_docx.find_record(draft) == {"docx": str(draft), "folder": str(deal), "candidates": [],
                                          "source": None, "ambiguous": True}
    r = fp_docx.build(draft, write(deal / "c.md", "## summary\nS.\n"), deal / "FP-v01.docx", mode="existing",
                      template=tpl)
    assert r["written"] == ["summary"] and "Approve (officer)." in fp_docx.to_markdown(deal / "FP-v01.docx")


def test_fields_in_a_normal_round_are_refused_when_the_first_heading_lost_its_style(tpl, tmp_path):
    v1, v2 = tmp_path / "v1.docx", tmp_path / "v2.docx"
    fp_docx.build(tpl, write(tmp_path / "c.md", FIRST), v1, mode="first", template=tpl)
    unstyle_headings(v1, "1. Summary")
    with pytest.raises(fp_docx.Refuse, match="record starts with 'summary'"):
        fp_docx.build(v1, write(tmp_path / "c2.md", "## market-risk\nM.\n"), v2, fields={"Amount": "EUR 15m"})
    assert not v2.exists()


def test_record_cannot_override_the_record_beside_the_base(tpl, tmp_path):
    v1, v2, v3 = (tmp_path / f"v{n}.docx" for n in (1, 2, 3))
    fp_docx.build(tpl, write(tmp_path / "c.md", FIRST), v1, mode="first", template=tpl)
    fp_docx.build(v1, write(tmp_path / "c2.md", "## summary\nNew summary.\n"), v2, protect=["market-risk"])
    before = v2.read_bytes(), (tmp_path / "v2.sections.json").read_bytes()
    with pytest.raises(fp_docx.Refuse, match="--record.*v2.sections.json.*exists"):
        fp_docx.build(v2, write(tmp_path / "c3.md", "## market-risk\nM.\n"), v3,
                      record=tmp_path / "v1.sections.json")
    assert not v3.exists() and not (tmp_path / "v3.sections.json").exists()
    assert before == (v2.read_bytes(), (tmp_path / "v2.sections.json").read_bytes())


# ---------- Task 2c: one set of incidental bookmarks for base and output ----------

SUM_RUN, FIG_RUN, REC_RUN = (f'<w:r><w:t xml:space="preserve">[{t}]</w:t></w:r>'
                             for t in ("Summarise the proposal", "Explain the figures", "State the recommendation"))


def test_hidden_bookmark_across_two_sections_lets_each_be_rewritten_without_duplicates(tpl, tmp_path):
    # X20: Word's copy marker starts in summary and ends in financial-analysis; redrafting summary orphans the end.
    v1, v2, v3 = (tmp_path / f"v{n}.docx" for n in (1, 2, 3))
    fp_docx.build(tpl, write(tmp_path / "c.md", FIRST), v1, mode="first", template=tpl)
    summary, revenue = (f'<w:r><w:t xml:space="preserve">{t}</w:t></w:r>'
                        for t in ("Cedar Foods seeks EUR 12m.", "Revenue rose 20%."))
    fp_docx.edit_text(v1, summary, f'<w:bookmarkStart w:id="9" w:name="{HLK}"/>{summary}')
    fp_docx.edit_text(v1, revenue, f'{revenue}<w:bookmarkEnd w:id="9"/>')
    assert fp_docx.build(v1, write(tmp_path / "c2.md", "## summary\nNew summary.\n"), v2)["written"] == ["summary"]
    r = fp_docx.build(v2, write(tmp_path / "c3.md", "## financial-analysis\nRevenue rose 25%.\n\n{{keep:1}}\n"), v3)
    assert r["written"] == ["financial-analysis"]
    md = fp_docx.to_markdown(v3)
    assert "Revenue rose 25%." in md and "Revenue rose 20%." not in md and "FY2025 (template)" in md


def test_link_into_another_section_does_not_block_rewriting_its_own_section(tpl, tmp_path):
    # X1: the summary links to a hidden bookmark in the recommendation; only the summary is redrafted.
    draft, v1, v2 = tmp_path / "draft.docx", tmp_path / "v1.docx", tmp_path / "v2.docx"
    shutil.copy(tpl, draft)
    fp_docx.edit_text(draft, REC_RUN, f'<w:bookmarkStart w:id="5" w:name="{HLK}"/><w:r><w:t>Officer recommendation.'
                      '</w:t></w:r><w:bookmarkEnd w:id="5"/>')
    fp_docx.edit_text(draft, SUM_RUN, '<w:r><w:t xml:space="preserve">Officer summary, </w:t></w:r>'
                      f'<w:hyperlink w:anchor="{HLK}"><w:r><w:t>see the recommendation</w:t></w:r></w:hyperlink>')
    r = fp_docx.build(draft, write(tmp_path / "c.md", "## summary\nNew summary.\n"), v1, mode="existing", template=tpl)
    assert r["written"] == ["summary"]
    r = fp_docx.build(v1, write(tmp_path / "c2.md", "## summary\nNewer summary.\n\n## recommendation\nDecline.\n"), v2,
                      template=tpl)
    assert r["written"] == ["summary"]
    assert r["frozen"] == {"recommendation": "not written by the assistant in the base version"}
    md = fp_docx.to_markdown(v2)
    assert "Newer summary." in md and "New summary." not in md and "Officer summary" not in md
    assert "Officer recommendation." in md and "Decline." not in md


@pytest.mark.parametrize("end_in", ["summary", "financial-analysis"])
def test_render_that_drops_a_referenced_bookmark_is_still_refused(tpl, tmp_path, monkeypatch, end_in):
    draft = tmp_path / "draft.docx"
    shutil.copy(tpl, draft)
    end = '<w:bookmarkEnd w:id="5"/>'
    fp_docx.edit_text(draft, SUM_RUN, f'<w:bookmarkStart w:id="5" w:name="{HLK}"/><w:r><w:t>Officer summary.</w:t></w:r>'
                      + (end if end_in == "summary" else ""))
    if end_in == "financial-analysis":       # the end survives in a section the build does not touch
        fp_docx.edit_text(draft, FIG_RUN, FIG_RUN + end)
    rec = "[State the recommendation]</w:t></w:r></w:p>"
    fp_docx.edit_text(draft, rec, rec + ANCHOR)   # the link that makes the bookmark functional
    monkeypatch.setattr(fp_docx, "render", lambda doc, *rest: fp_docx.Gen(doc).p("Lost the bookmark.", None))
    with pytest.raises(fp_docx.Refuse, match="fixed objects"):
        fp_docx.build(draft, write(tmp_path / "c.md", "## summary\nNew summary.\n"), tmp_path / "v1.docx",
                      mode="existing", template=tpl)
    assert not (tmp_path / "v1.docx").exists()
